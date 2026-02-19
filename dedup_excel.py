#!/usr/bin/env python3
"""
CCE 항목 중복 제거 + 플랫폼별 내용 필터링
1) 주요기반시설 가이드 항목의 '대상' 필드를 기반으로 해당 앱에만 할당
2) 점검 및 조치 사례에서 해당 플랫폼 섹션만 추출
3) 클라우드 가이드와의 중복항목을 병합
"""
import fitz
import re
import difflib
from collections import defaultdict, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = '/home/seongjin0526/cce_vuln_check'
CLOUD_PDF = f'{BASE_DIR}/클라우드 취약점 점검 가이드(2024).pdf'
MAIN_PDF = f'{BASE_DIR}/주요정보통신기반시설 기술적 취약점 분석·평가 방법 상세가이드.pdf'
OUTPUT_FILE = f'{BASE_DIR}/진단항목통합.xlsx'

CLOUD_OFFSET = 5

CLOUD_SECTIONS_ALL = [
    ('KVM', 7), ('Xenserver', 15), ('ESXi', 61), ('Hyper-V', 101),
    ('Server(Linux)', 119), ('Server(Windows)', 163),
    ('PC(Windows)', 225), ('PC(MAC)', 253), ('PC(Linux)', 283),
    ('MY-SQL', 299), ('MS-SQL', 311), ('Redis', 325), ('Elasticsearch', 335),
    ('MongoDB', 351), ('PostgresSQL', 363), ('Cubrid', 377), ('CouchDB', 389),
    ('SQLite', 407), ('Tibero', 417), ('InfluxDB', 429), ('Oracle', 441),
    ('Apache', 457), ('Nginx', 467), ('IIS', 477), ('Tomcat', 493),
    ('Docker', 505), ('Kubernetes(Master)', 543), ('Kubernetes(Worker)', 565),
    ('OpenStack', 579), ('PHP', 639), ('RabbitMQ', 647), ('Node.js', 659),
    ('Ceph', 671), ('Hadoop', 681), ('Network Device', 693),
    ('정보보호시스템', 721), ('스토리지', 737), ('BOSH(Director)', 745), ('BOSH(UAA)', 753),
]

CLOUD_TARGETS = {
    'KVM': 'KVM', 'Xenserver': 'Xenserver', 'ESXi': 'ESXi',
    'Linux': 'Server(Linux)', 'Windows': 'Server(Windows)',
    'MY-SQL': 'MY-SQL', 'MS-SQL': 'MS-SQL', 'Redis': 'Redis',
    'Elasticsearch': 'Elasticsearch', 'MongoDB': 'MongoDB',
    'PostgreSQL': 'PostgresSQL', 'Apache': 'Apache', 'Nginx': 'Nginx',
    'Tomcat': 'Tomcat', 'Docker': 'Docker',
    'K8s(Master)': 'Kubernetes(Master)', 'K8s(Worker)': 'Kubernetes(Worker)',
    'PHP': 'PHP', 'NodeJS': 'Node.js', 'Hadoop': 'Hadoop', 'Ceph': 'Ceph',
}

MAIN_CHAPTERS = {
    'Ch1_UNIX': {'start': 6, 'end': 171, 'pattern': r'(U-\d+)\s*\n\s*\(([상중하])\)'},
    'Ch2_Windows': {'start': 171, 'end': 270, 'pattern': r'(W-\d+)\s*\n\s*\(([상중하])\)'},
    'Ch3_Web': {'start': 270, 'end': 352, 'pattern': r'(WEB-\d+)\s*\n\s*\(([상중하])\)'},
    'Ch8_DBMS': {'start': 592, 'end': 669, 'pattern': r'(D-\d+)\s*\n\s*\(([상중하])\)'},
    'Ch11_Virtual': {'start': 786, 'end': 850, 'pattern': r'(HV-\d+)\s*\n\s*\(([상중하])\)'},
}

# Main guide chapter -> target apps + platform keywords for filtering
MAIN_TO_APPS = {
    'Ch1_UNIX': {
        'Linux': {
            'target_keywords': ['LINUX', 'Linux', 'Redhat', 'Debian', 'Ubuntu', 'CentOS'],
            'filter_detail': True,  # filter detail to only Linux sections
        },
    },
    'Ch2_Windows': {
        'Windows': {
            'target_keywords': ['Windows'],
            'filter_detail': False,  # all content is Windows
        },
    },
    'Ch3_Web': {
        'Apache': {
            'target_keywords': ['Apache', 'httpd'],
            'filter_detail': True,
        },
        'Nginx': {
            'target_keywords': ['Nginx', 'nginx'],
            'filter_detail': True,
        },
        'Tomcat': {
            'target_keywords': ['Tomcat', 'tomcat'],
            'filter_detail': True,
        },
    },
    'Ch8_DBMS': {
        'MY-SQL': {
            'target_keywords': ['MySQL', 'mysql', 'MY-SQL'],
            'filter_detail': True,
        },
        'MS-SQL': {
            'target_keywords': ['MSSQL', 'MS-SQL', 'SQL Server'],
            'filter_detail': True,
        },
        'PostgreSQL': {
            'target_keywords': ['PostgreSQL', 'postgres'],
            'filter_detail': True,
        },
    },
    'Ch11_Virtual': {
        'ESXi': {
            'target_keywords': ['VMware', 'ESXi', 'vCenter', 'vSphere'],
            'filter_detail': True,
        },
        'KVM': {
            'target_keywords': ['KVM'],
            'filter_detail': True,
        },
        'Xenserver': {
            'target_keywords': ['XenServer', 'Xen Server', 'Xenserver'],
            'filter_detail': True,
        },
    },
}

# Semantic match overrides for dedup
MATCH_OVERRIDES = {
    ('패스워드 최대 사용 기간', 'NTP 및 시각 동기화 설정'): False,
    ('패스워드 최대 사용 기간 설정', '계정 잠금 기간 설정'): False,
    ('암호 사용 기간 제한없음 제거', '하드디스크 기본 공유 제거'): False,
    ('최신 서비스팩 적용', '불필요한 서비스 제거'): False,
    ('SSH 데몬 빈암호 사용 인증 허용 제한', '가상화 장비 사용자 인증 강화'): False,
    ('SSL 시간 초과 구성 설정 확인', 'ESXi Shell 세션 종료 시간 설정'): False,
    ('불필요한 서비스 제거', '가상 머신의 불필요한 장치 제거'): False,
    ('FTP 비활성화', 'ESXi Shell 비활성화'): False,
    ('사용자 계정 관리', '가상화 장비 루트계정 관리'): False,
    ('NTP 시간 동기화 설정', '계정 잠금 임계값 설정'): False,
    ('불필요한 계정 제거', '가상 머신의 불필요한 장치 제거'): False,
    ('IP 접근 제한 설정', '가상머신의 장치 변경 제한 설정'): False,
    ('로깅 수준 설정', '비밀번호 관리정책 설정'): False,
    ('로그 파일 권한 설정', '계정 잠금 임계값 설정'): False,
    ('SU(Select User) 사용 제한', 'MOB'): False,
    ('로그인이 불필요한 계정 shell 제한', '가상 머신의 불필요한 장치 제거'): False,
    ('SSH(Secure Shell) 버전 취약점', 'ESXi Shell 비활성화'): False,
    ('Remote Shell 접근 제어', 'ESXi Shell 비활성화'): False,
    ('홈 디렉터리 쓰기 권한 관리', '로그 디렉터리 및 파일 권한 설정'): False,
    ('로그 파일 관리 및 주기적 백업', '비밀번호 파일 권한 관리'): False,
    ('환경 설정 파일 권한 관리', '비밀번호 파일 권한 관리'): False,
    ('Public schema 사용 제한', 'xp_cmdshell 사용 제한'): False,
    ('IP 접근 제한 설정', '비밀번호 재사용에 대한 제약 설정'): False,
    ('데이터 디렉터리 권한 설정', '데이터베이스의 자원 제한 기능을 TRUE'): False,
    ('Guest 계정 사용 제한', 'xp_cmdshell 사용 제한'): False,
    ('타 사용자에 권한 부여 옵션 제한', '비밀번호 재사용에 대한 제약 설정'): False,
    ('사용자 계정 정보 테이블 접근 권한', 'DB 사용자 계정을 개별적으로 부여'): False,
    ('취약한 패스워드 사용 제한', 'root 권한으로 서비스 구동 제한'): False,
    ('FTP 서비스 구동 점검', 'DNS 서비스 구동 점검'): False,
    # Positive overrides
    ('패스워드 최대 사용 기간', '비밀번호 관리정책 설정'): True,
    ('패스워드 파일 보호', '비밀번호 파일 보호'): True,
    ('패스워드 복잡도 설정', '비밀번호 관리정책 설정'): True,
    ('Anonymous FTP 비활성화', '공유 서비스에 대한 익명 접근 제한'): True,
    ('cron 파일 소유자 및 권한 설정', 'crontab 설정파일 권한'): True,
    ('로그의 정기적 검토 및 백업', '정책에 따른 시스템 로깅 설정'): True,
    ('Sendmail 버전 점검', '메일 서비스 버전 점검'): True,
    ('일반 사용자의 Sendmail 실행 방지', '일반 사용자의 메일 서비스 실행 방지'): True,
    ('RPC 서비스 확인', '불필요한 RPC 서비스 비활성화'): True,
    ('NFS 서비스 비활성화', '불필요한 NFS 서비스 비활성화'): True,
    ('automountd 제거', '불필요한 automountd 제거'): True,
    ('최신 서비스팩 적용', '주기적 보안 패치'): True,
    ('최신 Hot Fix 적용', '최신 Windows OS Build'): True,
    ('FTP 서비스 구동 점검', 'FTP 서비스 정보 노출'): True,
    ('NTP 시간 동기화 설정', 'NTP 및 시각 동기화 설정'): True,
    ('Default 계정 관리', '기본 계정의 비밀번호'): True,
    ('일반계정 root 권한 관리', '가상화 장비 계정 권한 관리'): True,
    ('SU 로그 설정', '시스템 주요 이벤트 로그'): True,
    ('보안패치 적용', '주기적 보안 패치'): True,
}


def clean_text(text):
    if not text:
        return ""
    text = re.sub(r'<<PAGE_\d+>>', '', text)
    text = re.sub(r'\d+\s*_\s*클라우드\s*취약점\s*점검\s*가이드', '', text)
    text = re.sub(r'\|\s*한국인터넷진흥원\s*\|', '', text)
    text = re.sub(r'\d{4}\s*주요정보통신기반시설.*?가이드', '', text)
    text = re.sub(r'\d{2}\.\s*(Unix|Windows|가상화|DBMS|웹).*?가이드', '', text)
    text = re.sub(r'2\.\s*보안가이드_\s*\d+', '', text)
    text = re.sub(r'\n\s*\d+\s*\n', '\n', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def get_cloud_page_range(section_name):
    for i, (name, start_page) in enumerate(CLOUD_SECTIONS_ALL):
        if name == section_name:
            pdf_start = start_page + CLOUD_OFFSET - 1
            if i + 1 < len(CLOUD_SECTIONS_ALL):
                pdf_end = CLOUD_SECTIONS_ALL[i + 1][1] + CLOUD_OFFSET - 1
            else:
                pdf_end = pdf_start + 20
            return pdf_start, pdf_end
    return None, None


def extract_pages_text(doc, start, end):
    parts = []
    for p in range(start, min(end, len(doc))):
        parts.append(f"<<PAGE_{p+1}>>\n{doc[p].get_text()}")
    return '\n'.join(parts)


def parse_cloud_section(doc, section_name):
    """Parse cloud guide section items"""
    start, end = get_cloud_page_range(section_name)
    if start is None:
        return []

    text = extract_pages_text(doc, start, end)
    items = []
    pages = text.split('<<PAGE_')
    current_item = None

    for page_block in pages:
        if not page_block.strip():
            continue
        page_lines = page_block.split('\n', 1)
        if len(page_lines) < 2:
            continue
        page_text = page_lines[1]

        if '항목설명' in page_text:
            if current_item:
                items.append(current_item)
            lines = page_text.split('\n')
            item_name = ""
            for j, line in enumerate(lines):
                if '항목설명' in line:
                    for k in range(j - 1, -1, -1):
                        candidate = lines[k].strip()
                        if candidate and not re.match(r'^\d+\s*_', candidate) and \
                           not re.match(r'^2\.\d+\.?$', candidate) and \
                           not re.match(r'^[0-9]+$', candidate) and \
                           candidate not in ['보안가이드', '클라우드 취약점 점검 가이드'] and \
                           not candidate.startswith('2. 보안가이드') and len(candidate) > 1:
                            item_name = candidate
                            break
                    break
            current_item = {'name': item_name, 'raw_text': page_text}
        elif current_item:
            current_item['raw_text'] += '\n' + page_text

    if current_item:
        items.append(current_item)

    # Parse fields
    parsed = []
    for item in items:
        raw = item['raw_text']
        good = bad = diagnosis = remediation = note = ''

        # Parse criteria block
        cb_match = re.search(r'(?:진단\s*\n?기준|기준)\s*\n(.*?)(?=진단\s*\n?방법)', raw, re.DOTALL)
        if cb_match:
            cb = cb_match.group(1)
            gm = re.search(r'양호\s*\n?(.*?)(?=취약)', cb, re.DOTALL)
            if gm: good = gm.group(1).strip()
            bm = re.search(r'취약\s*\n?(.*?)$', cb, re.DOTALL)
            if bm: bad = bm.group(1).strip()
        else:
            gm = re.search(r'양호\s*\n?(.*?)(?=취약)', raw, re.DOTALL)
            if gm: good = gm.group(1).strip()
            bm = re.search(r'취약\s*\n?(.*?)(?=진단\s*\n?방법)', raw, re.DOTALL)
            if bm: bad = bm.group(1).strip()

        dm = re.search(r'진단\s*\n?방법\s*\n(.*?)(?=조치\s*\n?방법)', raw, re.DOTALL)
        if dm: diagnosis = dm.group(1).strip()

        rm = re.search(r'조치\s*\n?방법\s*\n(.*?)(?=비고\s*\n|$)', raw, re.DOTALL)
        if rm: remediation = rm.group(1).strip()

        nm = re.search(r'비고\s*\n(.*?)$', raw, re.DOTALL)
        if nm: note = nm.group(1).strip()

        diag_full = ""
        if good: diag_full += f"[양호] {clean_text(good)}\n"
        if bad: diag_full += f"[취약] {clean_text(bad)}\n"
        if diagnosis: diag_full += f"\n[진단방법]\n{clean_text(diagnosis)}"
        if note: diag_full += f"\n\n[비고] {clean_text(note)}"

        parsed.append({
            'name': item['name'],
            'diagnosis': diag_full.strip(),
            'remediation': clean_text(remediation),
            'category': '',
        })

    return parsed


def parse_cloud_categories(doc, section_name):
    start, end = get_cloud_page_range(section_name)
    if start is None: return {}
    text = extract_pages_text(doc, start, min(start + 3, end))
    cats = {}
    lines = text.split('\n')
    in_cl = False
    cur_cat = ""
    for line in lines:
        line = line.strip()
        if '진단 체크리스트' in line or '진단 항목' in line:
            in_cl = True; continue
        if in_cl and line in ['구분', '진단 항목']: continue
        if in_cl:
            cm = re.match(r'^([가-힣])\.\s*(.+)', line)
            if cm: cur_cat = cm.group(2).strip(); continue
            if '항목설명' in line: break
            if not line or line.startswith('[') or line.startswith('총') or re.match(r'^\d+$', line): continue
            if len(line) > 2 and cur_cat:
                cats[line] = cur_cat
    return cats


def filter_detail_by_platform(detail_text, target_keywords):
    """Extract only platform-specific sections from detail text"""
    if not detail_text:
        return ""

    # Split into platform sections by "l PLATFORM" pattern
    # Pattern: starts with "l " followed by platform name
    sections = re.split(r'(?=\nl\s+)', detail_text)

    filtered = []
    generic_intro = ""

    for section in sections:
        section = section.strip()
        if not section:
            continue

        # Check if this section starts with a platform marker
        platform_match = re.match(r'^l\s+(.+?)(?:\n|$)', section)
        if platform_match:
            platform_name = platform_match.group(1).strip()
            # Check if any target keyword matches this platform
            is_match = any(kw.lower() in platform_name.lower() for kw in target_keywords)
            if is_match:
                filtered.append(section)
        else:
            # This is intro/generic text before platform sections
            generic_intro = section

    if filtered:
        result = generic_intro + '\n' + '\n'.join(filtered) if generic_intro else '\n'.join(filtered)
        return result.strip()
    elif generic_intro:
        return generic_intro.strip()
    else:
        # No platform sections found, return all text
        return detail_text


def check_target_applicability(target_field, target_keywords):
    """Check if this item applies to the target platform based on '대상' field"""
    if not target_field:
        return True  # no target info, assume applicable

    target_lower = target_field.lower()
    return any(kw.lower() in target_lower for kw in target_keywords)


def parse_main_chapter(doc, chapter_key):
    """Parse main guide chapter with full platform info"""
    config = MAIN_CHAPTERS[chapter_key]
    text = extract_pages_text(doc, config['start'], config['end'])
    items = []
    matches = list(re.finditer(config['pattern'], text))

    for i, match in enumerate(matches):
        code = match.group(1)
        severity = match.group(2)
        start_pos = match.start()
        end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        item_text = text[start_pos:end_pos]

        # Parse item name
        lines = item_text.split('\n')
        item_name = ''
        category_path = ''
        found_severity = False
        prefix = code.split('-')[0]

        for j, line in enumerate(lines):
            ls = line.strip()
            if re.match(rf'^{re.escape(prefix)}-\d+$', ls): continue
            if re.match(r'^\([상중하]\)$', ls): found_severity = True; continue
            if found_severity and '>' in ls and not item_name:
                category_path = ls; continue
            if found_severity and ls and category_path and not item_name:
                if ls.startswith('개요') or ls.startswith('점검') or ls.startswith('보안 위협') or ls.startswith('참고'):
                    continue
                item_name = ls; break

        # Extract target platforms
        target_field = ''
        tm = re.search(r'대상\s*\n(.*?)(?=판단\s*기준)', item_text, re.DOTALL)
        if tm:
            target_field = tm.group(1).strip()

        # Extract criteria
        good = bad = ''
        gm = re.search(r'양호\s*[：:]\s*(.*?)(?=취약\s*[：:])', item_text, re.DOTALL)
        if gm: good = gm.group(1).strip()
        bm = re.search(r'취약\s*[：:]\s*(.*?)(?=조치\s*방법)', item_text, re.DOTALL)
        if bm: bad = bm.group(1).strip()

        # Extract check content
        check = ''
        cm = re.search(r'점검\s*내용\s*\n?(.*?)(?=점검\s*목적)', item_text, re.DOTALL)
        if cm: check = cm.group(1).strip()

        # Extract remedy brief
        remedy_brief = ''
        rbm = re.search(r'조치\s*방법\s*\n(.*?)(?=조치\s*시\s*영향)', item_text, re.DOTALL)
        if rbm: remedy_brief = rbm.group(1).strip()

        # Extract impact
        impact = ''
        im = re.search(r'조치\s*시\s*영향\s*\n(.*?)(?=점검\s*및\s*조치\s*사례)', item_text, re.DOTALL)
        if im: impact = im.group(1).strip()

        # Extract detail (점검 및 조치 사례)
        detail = ''
        dm = re.search(r'점검\s*및\s*조치\s*사례\s*\n(.*?)$', item_text, re.DOTALL)
        if dm: detail = dm.group(1).strip()

        items.append({
            'code': code,
            'severity': severity,
            'name': item_name,
            'category': category_path,
            'target_field': target_field,
            'check_content': check,
            'good_criteria': good,
            'bad_criteria': bad,
            'remedy_brief': remedy_brief,
            'impact': impact,
            'detail': detail,
        })

    return items


def build_main_item_for_app(item, app_name, app_config):
    """Build a main guide item filtered for a specific app"""
    # Check if item applies to this app
    if item['target_field']:
        if not check_target_applicability(item['target_field'], app_config['target_keywords']):
            return None  # item doesn't apply to this app

    # Filter detail text
    filtered_detail = item['detail']
    if app_config.get('filter_detail') and filtered_detail:
        filtered_detail = filter_detail_by_platform(filtered_detail, app_config['target_keywords'])

    # Build diagnosis text
    diag_parts = []
    if item['check_content']:
        diag_parts.append(f"[점검내용] {clean_text(item['check_content'])}")
    if item['good_criteria']:
        diag_parts.append(f"[양호] {clean_text(item['good_criteria'])}")
    if item['bad_criteria']:
        diag_parts.append(f"[취약] {clean_text(item['bad_criteria'])}")
    if filtered_detail:
        diag_parts.append(f"\n[점검 및 조치 사례]\n{clean_text(filtered_detail)}")
    diagnosis = '\n'.join(diag_parts)

    # Build remediation text
    remed_parts = []
    if item['remedy_brief']:
        remed_parts.append(clean_text(item['remedy_brief']))
    if filtered_detail:
        remed_parts.append(f"\n[상세 조치 사례]\n{clean_text(filtered_detail)}")
    remediation = '\n'.join(remed_parts)

    return {
        'code': item['code'],
        'severity': item['severity'],
        'name': item['name'],
        'category': item['category'],
        'diagnosis': diagnosis.strip(),
        'remediation': remediation.strip(),
        'impact': clean_text(item['impact']) if item['impact'] else '',
    }


def is_semantic_match(cloud_name, main_name):
    cn = cloud_name.strip()
    mn = main_name.strip()

    # Check explicit overrides
    for (ck, mk), result in MATCH_OVERRIDES.items():
        if (ck in cn and mk in mn) or (mk in cn and ck in mn):
            return result

    ratio = difflib.SequenceMatcher(None, cn, mn).ratio()
    return ratio >= 0.55


def find_best_match(cloud_item, main_items, used):
    cn = cloud_item['name'].strip()
    best_ratio = 0
    best_idx = -1
    best_item = None

    for idx, mi in enumerate(main_items):
        if idx in used: continue
        mn = mi['name'].strip()
        ratio = difflib.SequenceMatcher(None, cn, mn).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best_idx = idx
            best_item = mi

    if best_item and is_semantic_match(cloud_item['name'], best_item['name']):
        return best_idx, best_item, best_ratio
    return -1, None, 0


def merge_items(ci, mi):
    code = f"{ci['code']} / {mi['code']}"
    severity = mi.get('severity', '-')
    name = ci['name']
    category = ci.get('category', '') or mi.get('category', '')

    diag_parts = []
    if ci.get('diagnosis'):
        diag_parts.append(f"[클라우드 가이드]\n{ci['diagnosis']}")
    if mi.get('diagnosis'):
        diag_parts.append(f"\n[주요기반시설 가이드]\n{mi['diagnosis']}")

    remed_parts = []
    if ci.get('remediation'):
        remed_parts.append(f"[클라우드 가이드]\n{ci['remediation']}")
    if mi.get('remediation'):
        remed_parts.append(f"\n[주요기반시설 가이드]\n{mi['remediation']}")

    return {
        'code': code,
        'severity': severity,
        'name': name,
        'category': category,
        'diagnosis': '\n'.join(diag_parts),
        'remediation': '\n'.join(remed_parts),
        'impact': mi.get('impact', ''),
        'source': '통합',
    }


def create_excel(all_data, output_path):
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "요약"

    hfont = Font(name='맑은 고딕', bold=True, size=11, color='FFFFFF')
    hfill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    halign = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col, (h, w) in enumerate(zip(
        ['진단대상', '클라우드 고유', '주요기반시설 고유', '통합(중복병합)', '합계'],
        [20, 18, 22, 22, 10]
    ), 1):
        c = ws_summary.cell(row=1, column=col, value=h)
        c.font = hfont; c.fill = hfill; c.alignment = halign
        ws_summary.column_dimensions[get_column_letter(col)].width = w

    ws = wb.create_sheet("CCE 항목 통합(중복제거)")
    cfont = Font(name='맑은 고딕', size=10)
    calign = Alignment(vertical='top', wrap_text=True)
    ccenter = Alignment(horizontal='center', vertical='top', wrap_text=True)
    mfill = PatternFill(start_color='DAEEF3', end_color='DAEEF3', fill_type='solid')
    clfill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    mnfill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    sfill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    headers = ['진단대상', '출처', '중요도', '항목번호', '항목명', '항목분류',
               '진단방법\n(판단기준 및 명령어)', '조치방법\n(조치 명령어)', '조치영향도']
    widths = [14, 12, 8, 18, 30, 18, 70, 70, 30]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hfont; c.fill = hfill; c.alignment = halign; c.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = 'A2'

    row = 2
    total = 0
    srow = 2
    gc, gm, gmerged = 0, 0, 0

    for app, data in all_data.items():
        cloud = data.get('cloud', [])
        main = data.get('main', [])

        merged_results = []
        cloud_only = []
        used_main = set()

        for ci in cloud:
            idx, mi, ratio = find_best_match(ci, main, used_main)
            if idx >= 0:
                merged_results.append(merge_items(ci, mi))
                used_main.add(idx)
            else:
                cloud_only.append(ci)

        main_only = [main[i] for i in range(len(main)) if i not in used_main]

        nc, nm, nmerge = len(cloud_only), len(main_only), len(merged_results)
        nt = nc + nm + nmerge
        gc += nc; gm += nm; gmerged += nmerge

        ws_summary.cell(row=srow, column=1, value=app)
        ws_summary.cell(row=srow, column=2, value=nc)
        ws_summary.cell(row=srow, column=3, value=nm)
        ws_summary.cell(row=srow, column=4, value=nmerge)
        ws_summary.cell(row=srow, column=5, value=nt)
        srow += 1

        # Section header
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        ws.cell(row=row, column=1,
                value=f"  {app} (통합:{nmerge}, 클라우드고유:{nc}, 주요기반시설고유:{nm}, 합계:{nt})").font = Font(name='맑은 고딕', bold=True, size=11)
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).fill = sfill
            ws.cell(row=row, column=c).border = border
        row += 1

        def write_row(item, source, fill):
            nonlocal row, total
            vals = [
                (app, ccenter), (source, ccenter),
                (item.get('severity', '-'), ccenter),
                (item.get('code', ''), ccenter),
                (item.get('name', ''), calign),
                (item.get('category', ''), calign),
                (item.get('diagnosis', ''), calign),
                (item.get('remediation', ''), calign),
                (item.get('impact', ''), calign),
            ]
            for col, (v, a) in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.font = cfont; c.alignment = a; c.border = border
                if col == 2: c.fill = fill
            row += 1; total += 1

        for item in merged_results:
            write_row(item, '통합', mfill)
        for item in cloud_only:
            write_row(item, '클라우드', clfill)
        for item in main_only:
            write_row(item, '주요기반시설', mnfill)

    # Summary total
    for col, val in enumerate([
        '합계', gc, gm, gmerged, gc + gm + gmerged
    ], 1):
        ws_summary.cell(row=srow, column=col, value=val).font = Font(bold=True)

    wb.save(output_path)
    print(f"\n{'='*60}")
    print(f"[완료] {output_path}")
    print(f"  통합(중복병합): {gmerged}건")
    print(f"  클라우드 고유:  {gc}건")
    print(f"  주요기반시설 고유: {gm}건")
    print(f"  총 항목: {total}건")


def main():
    print("=" * 60)
    print("CCE 항목 플랫폼 필터링 + 중복 제거")
    print("=" * 60)

    # 1. Extract cloud items
    print("\n[1/3] 클라우드 가이드 추출...")
    cdoc = fitz.open(CLOUD_PDF)
    cloud_data = {}
    for display_name, section_name in CLOUD_TARGETS.items():
        items = parse_cloud_section(cdoc, section_name)
        cats = parse_cloud_categories(cdoc, section_name)
        for item in items:
            item['category'] = cats.get(item['name'], '')
        cloud_data[display_name] = items
        print(f"  {display_name}: {len(items)}건")
    cdoc.close()

    # 2. Extract main items with platform filtering
    print("\n[2/3] 주요기반시설 가이드 추출 + 플랫폼 필터링...")
    mdoc = fitz.open(MAIN_PDF)
    main_raw = {}
    for ch_key in MAIN_CHAPTERS:
        main_raw[ch_key] = parse_main_chapter(mdoc, ch_key)
        print(f"  {ch_key}: {len(main_raw[ch_key])}건 원본 추출")
    mdoc.close()

    # Build per-app main items with filtering
    main_data = defaultdict(list)
    for ch_key, app_configs in MAIN_TO_APPS.items():
        raw_items = main_raw.get(ch_key, [])
        for app_name, app_config in app_configs.items():
            count = 0
            for raw_item in raw_items:
                filtered = build_main_item_for_app(raw_item, app_name, app_config)
                if filtered:
                    main_data[app_name].append(filtered)
                    count += 1
            print(f"  {ch_key} → {app_name}: {count}건 적용")

    # 3. Combine and dedup
    print("\n[3/3] 통합 및 중복 제거...")
    app_order = [
        'KVM', 'Xenserver', 'ESXi', 'Linux', 'Windows',
        'MY-SQL', 'MS-SQL', 'Redis', 'Elasticsearch', 'MongoDB', 'PostgreSQL',
        'Apache', 'Nginx', 'Tomcat', 'Docker', 'K8s(Master)', 'K8s(Worker)',
        'PHP', 'NodeJS', 'Hadoop', 'Ceph',
    ]

    all_data = OrderedDict()
    for app in app_order:
        ci = cloud_data.get(app, [])
        mi = main_data.get(app, [])

        # Assign codes to cloud items
        for idx, item in enumerate(ci, 1):
            code_app = app.replace('(', '').replace(')', '').replace(' ', '')
            item['code'] = f"CLD-{code_app}-{idx:02d}"
            item['severity'] = '-'

        all_data[app] = {'cloud': ci, 'main': mi}

        if ci and mi:
            used = set()
            matched = 0
            for c in ci:
                for j, m in enumerate(mi):
                    if j in used: continue
                    if is_semantic_match(c['name'], m['name']):
                        matched += 1; used.add(j); break
            print(f"  {app}: 클라우드 {len(ci)} + 주요기반시설 {len(mi)} → 중복 {matched}건 병합 → {len(ci) + len(mi) - matched}건")
        else:
            print(f"  {app}: {len(ci) + len(mi)}건")

    create_excel(all_data, OUTPUT_FILE)


if __name__ == '__main__':
    main()
