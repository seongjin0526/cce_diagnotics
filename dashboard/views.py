from __future__ import annotations

import io
import json
from collections import Counter

from flask import (
    Blueprint,
    Response,
    abort,
    current_app,
    flash,
    g,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from openpyxl import Workbook
from werkzeug.security import check_password_hash

from .auth import login_required, security_required
from .catalog import (
    deserialize_frameworks,
    framework_options,
    get_app_definition,
    get_control,
    get_controls_for_app,
    host_support_for_app,
    list_app_definitions,
    list_controls,
)
from .db import execute, fetch_all, fetch_one, utcnow
from .jobs import launch_assessment, launch_discovery


bp = Blueprint("main", __name__)

ALLOWED_TRANSPORTS = {"local", "ssh", "compose"}
ALLOWED_FINAL_STATUSES = {"양호", "취약", "N/A", "수동점검", "예외"}
ALLOWED_APPROVAL_DECISIONS = {"approved", "rejected"}


def effective_status(row) -> str:
    return row["final_status"] or row["script_status"]


def latest_exception_requests(result_ids: list[int]) -> dict[int, dict]:
    if not result_ids:
        return {}
    placeholders = ",".join(["?"] * len(result_ids))
    rows = fetch_all(
        f"""
        SELECT er.*
        FROM exception_requests er
        JOIN (
            SELECT assessment_result_id, MAX(id) AS max_id
            FROM exception_requests
            WHERE assessment_result_id IN ({placeholders})
            GROUP BY assessment_result_id
        ) latest
          ON latest.max_id = er.id
        """,
        tuple(result_ids),
    )
    return {row["assessment_result_id"]: row for row in rows}


@bp.app_template_filter("status_badge")
def status_badge(status: str) -> str:
    mapping = {
        "양호": "status-good",
        "취약": "status-bad",
        "예외": "status-exception",
        "N/A": "status-na",
        "수동점검": "status-manual",
        "pending": "status-pending",
        "approved": "status-good",
        "rejected": "status-bad",
    }
    return mapping.get(status, "status-default")


@bp.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = fetch_one("SELECT * FROM users WHERE username = ?", (username,))
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            return redirect(url_for("main.index"))
        flash("로그인 정보가 올바르지 않습니다.", "error")
    return render_template("login.html")


@bp.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("main.login"))


@bp.route("/")
@login_required
def index():
    summary = {
        "hosts": fetch_one("SELECT COUNT(*) AS count FROM hosts")["count"],
        "discoveries": fetch_one("SELECT COUNT(*) AS count FROM discovery_runs")["count"],
        "assessments": fetch_one("SELECT COUNT(*) AS count FROM assessment_runs")["count"],
        "pending_approvals": fetch_one(
            "SELECT COUNT(*) AS count FROM exception_requests WHERE status = 'pending'"
        )["count"],
    }
    hosts = fetch_all("SELECT * FROM hosts ORDER BY id DESC")
    recent_runs = fetch_all(
        """
        SELECT ar.*, h.name AS host_name
        FROM assessment_runs ar
        JOIN hosts h ON h.id = ar.host_id
        ORDER BY ar.id DESC
        LIMIT 10
        """
    )
    return render_template("index.html", summary=summary, hosts=hosts, recent_runs=recent_runs)


@bp.route("/hosts", methods=["GET", "POST"])
@login_required
def hosts():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        address = request.form.get("address", "").strip()
        transport = request.form.get("transport", "").strip()
        if not name or not address:
            flash("호스트 이름과 주소를 입력해야 합니다.", "error")
            return redirect(url_for("main.hosts"))
        if transport not in ALLOWED_TRANSPORTS:
            flash("지원하지 않는 전송방식입니다.", "error")
            return redirect(url_for("main.hosts"))
        try:
            port = int(request.form.get("port") or 22)
        except ValueError:
            flash("포트는 숫자로 입력해야 합니다.", "error")
            return redirect(url_for("main.hosts"))
        if port < 1 or port > 65535:
            flash("포트는 1부터 65535 사이여야 합니다.", "error")
            return redirect(url_for("main.hosts"))

        cursor = execute(
            """
            INSERT INTO hosts (name, address, port, remote_user, transport, shell_type, use_sudo, notes, created_by, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                name,
                address,
                port,
                request.form.get("remote_user", "").strip() or None,
                transport,
                "posix",
                1 if request.form.get("use_sudo") else 0,
                request.form.get("notes", "").strip(),
                g.user["id"],
                utcnow(),
            ),
        )
        flash("호스트를 등록했습니다.", "success")
        return redirect(url_for("main.host_detail", host_id=cursor.lastrowid))

    rows = fetch_all("SELECT * FROM hosts ORDER BY id DESC")
    return render_template("hosts.html", hosts=rows)


@bp.route("/hosts/<int:host_id>")
@login_required
def host_detail(host_id: int):
    host = fetch_one("SELECT * FROM hosts WHERE id = ?", (host_id,))
    if host is None:
        abort(404)

    latest_discovery = fetch_one(
        "SELECT * FROM discovery_runs WHERE host_id = ? ORDER BY id DESC LIMIT 1", (host_id,)
    )
    discovery_results = []
    if latest_discovery:
        discovery_results = fetch_all(
            "SELECT * FROM discovery_results WHERE run_id = ? ORDER BY app_key", (latest_discovery["id"],)
        )

    assessment_runs = fetch_all(
        "SELECT * FROM assessment_runs WHERE host_id = ? ORDER BY id DESC LIMIT 20", (host_id,)
    )
    settings_rows = fetch_all(
        "SELECT * FROM host_app_settings WHERE host_id = ? ORDER BY app_key", (host_id,)
    )
    settings_map = {row["app_key"]: row for row in settings_rows}
    controls_count = {app.key: len(get_controls_for_app(app.key)) for app in list_app_definitions()}
    support_matrix = []
    for app in list_app_definitions():
        support = host_support_for_app(host["shell_type"], host["transport"], app.key)
        support_matrix.append(
            {
                "app_key": app.display_name,
                "supported": support["supported"],
                "reason": support["reason"],
                "controls_count": controls_count[app.key],
            }
        )
    app_settings_cards = []
    for app in list_app_definitions():
        if not app.path_settings:
            continue
        row = settings_map.get(app.key)
        detected_paths = json.loads(row["detected_paths_json"]) if row and row["detected_paths_json"] else {}
        manual_paths = json.loads(row["manual_paths_json"]) if row and row["manual_paths_json"] else {}
        support = host_support_for_app(host["shell_type"], host["transport"], app.key)
        app_settings_cards.append(
            {
                "app": app,
                "detected_paths": detected_paths,
                "manual_paths": manual_paths,
                "last_detected_at": row["last_detected_at"] if row else None,
                "supported": support["supported"],
                "support_reason": support["reason"],
            }
        )
    auto_refresh = bool(
        (latest_discovery and latest_discovery["status"] in {"queued", "running"})
        or any(run["status"] in {"queued", "running"} for run in assessment_runs)
    )
    return render_template(
        "host_detail.html",
        host=host,
        latest_discovery=latest_discovery,
        discovery_results=discovery_results,
        assessment_runs=assessment_runs,
        controls_count=controls_count,
        support_matrix=support_matrix,
        app_settings_cards=app_settings_cards,
        auto_refresh=auto_refresh,
        app_definitions=list_app_definitions(),
    )


@bp.route("/hosts/<int:host_id>/discover", methods=["POST"])
@login_required
def run_discovery(host_id: int):
    host = fetch_one("SELECT * FROM hosts WHERE id = ?", (host_id,))
    if host is None:
        abort(404)

    active_run = fetch_one(
        """
        SELECT id FROM discovery_runs
        WHERE host_id = ? AND status IN ('queued', 'running')
        ORDER BY id DESC LIMIT 1
        """,
        (host_id,),
    )
    if active_run:
        flash("이미 실행 중인 탐지 작업이 있습니다.", "error")
        return redirect(url_for("main.host_detail", host_id=host_id))
    run_cursor = execute(
        """
        INSERT INTO discovery_runs (host_id, initiated_by, started_at, status)
        VALUES (?, ?, ?, ?)
        """,
        (host_id, g.user["id"], utcnow(), "queued"),
    )
    run_id = run_cursor.lastrowid
    execute(
        "UPDATE hosts SET last_discovery_status = ?, last_discovery_at = ? WHERE id = ?",
        ("queued", utcnow(), host_id),
    )
    launch_discovery(current_app._get_current_object(), host_id, run_id)
    flash("호스트 애플리케이션 탐지 작업을 큐에 등록했습니다.", "success")
    return redirect(url_for("main.host_detail", host_id=host_id))


@bp.route("/hosts/<int:host_id>/assess/<app_key>", methods=["POST"])
@login_required
def start_assessment(host_id: int, app_key: str):
    host = fetch_one("SELECT * FROM hosts WHERE id = ?", (host_id,))
    if host is None:
        abort(404)

    app = get_app_definition(app_key)
    if app is None:
        flash("알 수 없는 애플리케이션입니다.", "error")
        return redirect(url_for("main.host_detail", host_id=host_id))
    support = host_support_for_app(host["shell_type"], host["transport"], app_key)
    if not support["supported"]:
        flash(f"현재 호스트에서는 {app.display_name} 진단을 실행할 수 없습니다: {support['reason']}", "error")
        return redirect(url_for("main.host_detail", host_id=host_id))
    active_run = fetch_one(
        """
        SELECT id FROM assessment_runs
        WHERE host_id = ? AND app_key = ? AND status IN ('queued', 'running')
        ORDER BY id DESC LIMIT 1
        """,
        (host_id, app_key),
    )
    if active_run:
        flash("이미 실행 중인 동일 애플리케이션 진단이 있습니다.", "error")
        return redirect(url_for("main.host_detail", host_id=host_id))
    run_cursor = execute(
        """
        INSERT INTO assessment_runs (host_id, app_key, initiated_by, started_at, status, framework_scope)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (host_id, app_key, g.user["id"], utcnow(), "queued", "공공CSAP + ISMS-P"),
    )
    run_id = run_cursor.lastrowid
    launch_assessment(current_app._get_current_object(), host_id, run_id, app_key)
    flash(f"{app.display_name} 진단 작업을 큐에 등록했습니다.", "success")
    return redirect(url_for("main.run_detail", run_id=run_id))


@bp.route("/hosts/<int:host_id>/settings/<app_key>", methods=["POST"])
@login_required
def save_host_app_settings(host_id: int, app_key: str):
    host = fetch_one("SELECT * FROM hosts WHERE id = ?", (host_id,))
    app = get_app_definition(app_key)
    if host is None or app is None:
        flash("호스트 또는 애플리케이션 정보를 찾을 수 없습니다.", "error")
        return redirect(url_for("main.hosts"))

    manual_paths = {}
    for setting in app.path_settings:
        value = request.form.get(setting.env_var, "").strip()
        if value:
            manual_paths[setting.env_var] = value

    existing = fetch_one(
        "SELECT detected_paths_json, last_detected_at FROM host_app_settings WHERE host_id = ? AND app_key = ?",
        (host_id, app_key),
    )
    execute(
        """
        INSERT INTO host_app_settings
            (host_id, app_key, detected_paths_json, manual_paths_json, last_detected_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(host_id, app_key) DO UPDATE SET
            manual_paths_json = excluded.manual_paths_json,
            updated_at = excluded.updated_at
        """,
        (
            host_id,
            app_key,
            existing["detected_paths_json"] if existing else "{}",
            json.dumps(manual_paths, ensure_ascii=False),
            existing["last_detected_at"] if existing else None,
            utcnow(),
        ),
    )
    flash(f"{app.display_name} 경로 오버라이드를 저장했습니다.", "success")
    return redirect(url_for("main.host_detail", host_id=host_id, _anchor="paths"))


def _filtered_run_results(run_id: int):
    rows = fetch_all("SELECT * FROM assessment_results WHERE run_id = ? ORDER BY id", (run_id,))
    request_map = latest_exception_requests([row["id"] for row in rows])
    filtered = []
    framework_filter = request.args.get("framework", "").strip()
    source_filter = request.args.get("source", "").strip()
    status_filter = request.args.get("status", "").strip()
    query = request.args.get("q", "").strip().lower()

    for row in rows:
        item = dict(row)
        item["frameworks"] = deserialize_frameworks(row["framework_tags"])
        item["effective_status"] = effective_status(row)
        item["request"] = request_map.get(row["id"])

        if framework_filter and framework_filter not in item["frameworks"]:
            continue
        if source_filter and source_filter != (item["source"] or ""):
            continue
        if status_filter and status_filter != item["effective_status"]:
            continue
        if query:
            haystack = " ".join(
                [
                    item.get("code", ""),
                    item.get("title", ""),
                    item.get("category", ""),
                    item.get("detail", ""),
                    item.get("remediation", ""),
                ]
            ).lower()
            if query not in haystack:
                continue
        filtered.append(item)
    return filtered


@bp.route("/runs/<int:run_id>")
@login_required
def run_detail(run_id: int):
    run = fetch_one(
        """
        SELECT ar.*, h.name AS host_name
        FROM assessment_runs ar
        JOIN hosts h ON h.id = ar.host_id
        WHERE ar.id = ?
        """,
        (run_id,),
    )
    if run is None:
        abort(404)

    results = _filtered_run_results(run_id)
    counts = Counter(item["effective_status"] for item in results)
    return render_template(
        "run_detail.html",
        run=run,
        results=results,
        counts=counts,
        auto_refresh=run["status"] in {"queued", "running"},
        framework_options=framework_options(),
    )


@bp.route("/runs/<int:run_id>/export")
@login_required
def export_run(run_id: int):
    run = fetch_one("SELECT * FROM assessment_runs WHERE id = ?", (run_id,))
    if run is None:
        abort(404)

    results = _filtered_run_results(run_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "진단결과"
    headers = [
        "코드",
        "항목명",
        "카테고리",
        "원본판정",
        "최종판정",
        "출처",
        "프레임워크",
        "판단근거",
        "수행명령",
        "현재상태",
        "조치방법",
    ]
    sheet.append(headers)
    for item in results:
        sheet.append(
            [
                item["code"],
                item["title"],
                item["category"],
                item["script_status"],
                item["effective_status"],
                item["source"],
                ", ".join(item["frameworks"]),
                item["detail"],
                item["command"],
                item["current_state"],
                item["remediation"],
            ]
        )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"assessment_run_{run['id']}.xlsx"
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@bp.route("/controls")
@login_required
def controls():
    app_filter = request.args.get("app", "").strip()
    framework_filter = request.args.get("framework", "").strip()
    query = request.args.get("q", "").strip().lower()

    rows = list_controls(app_filter or None)
    filtered = []
    for row in rows:
        if framework_filter and framework_filter not in row["frameworks"]:
            continue
        if query:
            haystack = " ".join(
                [row["target"], row["code"], row["title"], row["category"], row["diagnosis"], row["remediation"]]
            ).lower()
            if query not in haystack:
                continue
        filtered.append(row)

    return render_template(
        "controls.html",
        controls=filtered,
        app_definitions=list_app_definitions(),
        framework_options=framework_options(),
    )


@bp.route("/results/<int:result_id>/exception", methods=["POST"])
@login_required
def request_exception(result_id: int):
    reason = request.form.get("reason", "").strip()
    row = fetch_one("SELECT * FROM assessment_results WHERE id = ?", (result_id,))
    if row is None:
        abort(404)

    if not reason:
        flash("예외 사유를 입력해야 합니다.", "error")
        return redirect(url_for("main.run_detail", run_id=row["run_id"]))

    execute(
        """
        INSERT INTO exception_requests (assessment_result_id, requested_by, requested_at, status, reason)
        VALUES (?, ?, ?, 'pending', ?)
        """,
        (result_id, g.user["id"], utcnow(), reason),
    )
    flash("예외 승인을 요청했습니다.", "success")
    return redirect(url_for("main.run_detail", run_id=row["run_id"]))


@bp.route("/results/<int:result_id>/override", methods=["POST"])
@security_required
def override_result(result_id: int):
    final_status = request.form.get("final_status", "").strip() or None
    note = request.form.get("status_note", "").strip() or None
    row = fetch_one("SELECT * FROM assessment_results WHERE id = ?", (result_id,))
    if row is None:
        abort(404)
    if final_status is not None and final_status not in ALLOWED_FINAL_STATUSES:
        flash("지원하지 않는 최종 판정입니다.", "error")
        return redirect(url_for("main.run_detail", run_id=row["run_id"]))

    if final_status == "예외":
        request_row = fetch_one(
            """
            SELECT * FROM exception_requests
            WHERE assessment_result_id = ? AND status = 'approved'
            ORDER BY id DESC LIMIT 1
            """,
            (result_id,),
        )
        if request_row is None:
            flash("승인된 예외 요청이 없어서 예외 판정을 적용할 수 없습니다.", "error")
            return redirect(url_for("main.run_detail", run_id=row["run_id"]))

    execute(
        "UPDATE assessment_results SET final_status = ?, status_note = ? WHERE id = ?",
        (final_status, note, result_id),
    )
    flash("최종 판정을 저장했습니다.", "success")
    return redirect(url_for("main.run_detail", run_id=row["run_id"]))


@bp.route("/approvals")
@security_required
def approvals():
    rows = fetch_all(
        """
        SELECT er.*, ar.run_id, ar.code, ar.title, h.name AS host_name, u.username AS requester_name
        FROM exception_requests er
        JOIN assessment_results ar ON ar.id = er.assessment_result_id
        JOIN assessment_runs run ON run.id = ar.run_id
        JOIN hosts h ON h.id = run.host_id
        JOIN users u ON u.id = er.requested_by
        ORDER BY er.status = 'pending' DESC, er.id DESC
        """
    )
    return render_template("approvals.html", requests=rows)


@bp.route("/approvals/<int:request_id>/decision", methods=["POST"])
@security_required
def approval_decision(request_id: int):
    decision = request.form.get("decision", "").strip()
    comment = request.form.get("approver_comment", "").strip()
    request_row = fetch_one("SELECT * FROM exception_requests WHERE id = ?", (request_id,))
    if request_row is None:
        abort(404)
    if request_row["status"] != "pending":
        flash("이미 처리된 예외 요청입니다.", "error")
        return redirect(url_for("main.approvals"))
    if decision not in ALLOWED_APPROVAL_DECISIONS:
        flash("지원하지 않는 승인 처리 값입니다.", "error")
        return redirect(url_for("main.approvals"))

    execute(
        """
        UPDATE exception_requests
           SET status = ?, approver_id = ?, decided_at = ?, approver_comment = ?
         WHERE id = ?
        """,
        (decision, g.user["id"], utcnow(), comment, request_id),
    )
    if decision == "approved":
        execute(
            "UPDATE assessment_results SET final_status = '예외', status_note = ? WHERE id = ?",
            (comment or "보안담당자 승인", request_row["assessment_result_id"]),
        )
    flash("승인 처리 결과를 저장했습니다.", "success")
    return redirect(url_for("main.approvals"))
