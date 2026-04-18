from __future__ import annotations

import json
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from code_scheme import control_code_aliases, to_preferred_code_text
from project_paths import MERGED_ITEMS_XLSX, REPO_ROOT, SCRIPTS_DIR


@dataclass(frozen=True)
class PathSetting:
    env_var: str
    label: str
    placeholder: str
    help_text: str
    probe_command: str


@dataclass(frozen=True)
class ApplicationDefinition:
    key: str
    display_name: str
    script_path: Path
    shell: str
    shell_family: str
    probe_rules: tuple[tuple[str, str], ...]
    path_settings: tuple[PathSetting, ...] = ()


def _rule(label: str, command: str) -> tuple[str, str]:
    return (label, command)


def _path(
    env_var: str,
    label: str,
    placeholder: str,
    help_text: str,
    probe_command: str,
) -> PathSetting:
    return PathSetting(
        env_var=env_var,
        label=label,
        placeholder=placeholder,
        help_text=help_text,
        probe_command=probe_command.strip(),
    )


APP_DEFINITIONS: OrderedDict[str, ApplicationDefinition] = OrderedDict(
    {
        "Linux": ApplicationDefinition(
            key="Linux",
            display_name="Linux",
            script_path=REPO_ROOT / "linux_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("os-release", "test -f /etc/os-release"),
                _rule("systemctl", "command -v systemctl >/dev/null 2>&1"),
            ),
        ),
        "Windows": ApplicationDefinition(
            key="Windows",
            display_name="Windows",
            script_path=SCRIPTS_DIR / "windows_cce_check.ps1",
            shell="powershell",
            shell_family="powershell",
            probe_rules=(),
        ),
        "KVM": ApplicationDefinition(
            key="KVM",
            display_name="KVM",
            script_path=SCRIPTS_DIR / "kvm_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("virsh", "command -v virsh >/dev/null 2>&1"),
                _rule("libvirtd", "ps -ef | grep -E '[l]ibvirtd|[v]irtqemud' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "LIBVIRT_CONF",
                    "libvirt 설정 디렉터리",
                    "/etc/libvirt",
                    "디렉터리 경로를 입력합니다.",
                    """
                    if [ -d /etc/libvirt ]; then
                        printf '%s' /etc/libvirt
                    fi
                    """,
                ),
            ),
        ),
        "Xenserver": ApplicationDefinition(
            key="Xenserver",
            display_name="XenServer",
            script_path=SCRIPTS_DIR / "xenserver_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("xe", "command -v xe >/dev/null 2>&1"),
                _rule("xapi", "ps -ef | grep -E '[x]api' >/dev/null 2>&1"),
            ),
        ),
        "ESXi": ApplicationDefinition(
            key="ESXi",
            display_name="ESXi",
            script_path=SCRIPTS_DIR / "esxi_cce_check.sh",
            shell="sh",
            shell_family="posix",
            probe_rules=(
                _rule("esxcli", "command -v esxcli >/dev/null 2>&1"),
                _rule("vmware-release", "test -f /etc/vmware-release"),
            ),
        ),
        "MY-SQL": ApplicationDefinition(
            key="MY-SQL",
            display_name="MySQL",
            script_path=SCRIPTS_DIR / "mysql_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("mysql", "command -v mysql >/dev/null 2>&1"),
                _rule("mysqld", "ps -ef | grep -E '[m]ysqld' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "MYSQL_CONF",
                    "MySQL 설정 파일",
                    "/etc/mysql/my.cnf",
                    "mysqld가 참조하는 메인 설정 파일 경로입니다.",
                    """
                    defaults_file=$(ps -ef 2>/dev/null | grep '[m]ysqld' | sed -n 's/.*--defaults-file=\\([^ ]*\\).*/\\1/p' | head -1)
                    if [ -n "$defaults_file" ] && [ -f "$defaults_file" ]; then
                        printf '%s' "$defaults_file"
                    else
                        for f in /etc/my.cnf /etc/mysql/my.cnf /etc/mysql/mysql.conf.d/mysqld.cnf "$HOME/.my.cnf" /usr/local/mysql/my.cnf; do
                            if [ -f "$f" ]; then
                                printf '%s' "$f"
                                break
                            fi
                        done
                    fi
                    """,
                ),
            ),
        ),
        "MS-SQL": ApplicationDefinition(
            key="MS-SQL",
            display_name="MSSQL",
            script_path=SCRIPTS_DIR / "mssql_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("sqlcmd", "command -v sqlcmd >/dev/null 2>&1"),
                _rule("sqlservr", "ps -ef | grep -E '[s]qlservr' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "MSSQL_CONF",
                    "MSSQL 설정 파일",
                    "/var/opt/mssql/mssql.conf",
                    "Linux SQL Server의 mssql.conf 경로입니다.",
                    """
                    for f in /var/opt/mssql/mssql.conf /opt/mssql/lib/mssql-conf/mssql.conf; do
                        if [ -f "$f" ]; then
                            printf '%s' "$f"
                            break
                        fi
                    done
                    """,
                ),
            ),
        ),
        "PostgreSQL": ApplicationDefinition(
            key="PostgreSQL",
            display_name="PostgreSQL",
            script_path=SCRIPTS_DIR / "postgresql_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("psql", "command -v psql >/dev/null 2>&1"),
                _rule("postgres", "ps -ef | grep -E '[p]ostgres' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "PG_DATA",
                    "PostgreSQL 데이터 디렉터리",
                    "/var/lib/postgresql/16/main",
                    "postgres 프로세스의 -D 값 또는 실제 데이터 디렉터리입니다.",
                    """
                    pg_data=''
                    pg_proc=$(ps -ef 2>/dev/null | grep '[p]ostgres.*-D' | head -1)
                    if [ -n "$pg_proc" ]; then
                        pg_data=$(printf '%s' "$pg_proc" | sed -n 's/.*-D[[:space:]]*\\([^ ]*\\).*/\\1/p')
                    fi
                    if [ -z "$pg_data" ]; then
                        pg_config_bin=$(command -v pg_config 2>/dev/null)
                        if [ -n "$pg_config_bin" ]; then
                            sharedir=$("$pg_config_bin" --sharedir 2>/dev/null)
                            if [ -n "$sharedir" ] && [ -d "$(dirname "$sharedir")/data" ]; then
                                pg_data=$(dirname "$sharedir")/data
                            fi
                        fi
                    fi
                    if [ -z "$pg_data" ]; then
                        for d in /var/lib/postgresql/*/main /var/lib/pgsql/*/data /var/lib/pgsql/data /usr/local/pgsql/data; do
                            if [ -d "$d" ]; then
                                pg_data="$d"
                                break
                            fi
                        done
                    fi
                    if [ -n "$pg_data" ]; then
                        printf '%s' "$pg_data"
                    fi
                    """,
                ),
                _path(
                    "PG_CONF",
                    "PostgreSQL 설정 파일",
                    "/etc/postgresql/16/main/postgresql.conf",
                    "postgresql.conf 전체 경로입니다.",
                    """
                    pg_conf=''
                    pg_data=''
                    pg_proc=$(ps -ef 2>/dev/null | grep '[p]ostgres.*-D' | head -1)
                    if [ -n "$pg_proc" ]; then
                        pg_data=$(printf '%s' "$pg_proc" | sed -n 's/.*-D[[:space:]]*\\([^ ]*\\).*/\\1/p')
                    fi
                    if [ -n "$pg_data" ] && [ -f "$pg_data/postgresql.conf" ]; then
                        pg_conf="$pg_data/postgresql.conf"
                    fi
                    if [ -z "$pg_conf" ]; then
                        for f in /etc/postgresql/*/main/postgresql.conf; do
                            if [ -f "$f" ]; then
                                pg_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$pg_conf" ]; then
                        printf '%s' "$pg_conf"
                    fi
                    """,
                ),
            ),
        ),
        "Redis": ApplicationDefinition(
            key="Redis",
            display_name="Redis",
            script_path=SCRIPTS_DIR / "redis_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("redis-cli", "command -v redis-cli >/dev/null 2>&1"),
                _rule("redis-server", "ps -ef | grep -E '[r]edis-server' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "REDIS_CONF",
                    "Redis 설정 파일",
                    "/etc/redis/redis.conf",
                    "redis-server가 참조하는 설정 파일 경로입니다.",
                    """
                    redis_proc=$(ps -ef 2>/dev/null | grep '[r]edis-server' | head -1)
                    if [ -n "$redis_proc" ]; then
                        conf_from_proc=$(printf '%s' "$redis_proc" | grep -oE '[^ ]+redis\\.conf' | head -1)
                        if [ -n "$conf_from_proc" ] && [ -f "$conf_from_proc" ]; then
                            printf '%s' "$conf_from_proc"
                        fi
                    fi
                    if [ ! -f "$conf_from_proc" ]; then
                        for f in /etc/redis/redis.conf /etc/redis.conf /etc/redis/6379.conf /usr/local/etc/redis.conf; do
                            if [ -f "$f" ]; then
                                printf '%s' "$f"
                                break
                            fi
                        done
                    fi
                    """,
                ),
            ),
        ),
        "Elasticsearch": ApplicationDefinition(
            key="Elasticsearch",
            display_name="Elasticsearch",
            script_path=SCRIPTS_DIR / "elasticsearch_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("elasticsearch", "ps -ef | grep -E '[e]lasticsearch' >/dev/null 2>&1"),
                _rule("es-data", "test -d /etc/elasticsearch || test -d /usr/share/elasticsearch"),
            ),
            path_settings=(
                _path(
                    "ES_CONF",
                    "Elasticsearch 설정 파일",
                    "/etc/elasticsearch/elasticsearch.yml",
                    "elasticsearch.yml 전체 경로입니다.",
                    """
                    es_conf=''
                    es_proc=$(ps -ef 2>/dev/null | grep '[e]lasticsearch' | grep -v grep | head -1)
                    if [ -n "$es_proc" ]; then
                        conf_from_proc=$(printf '%s' "$es_proc" | sed -n 's/.*-Epath\\.conf=\\([^ ]*\\).*/\\1/p')
                        if [ -n "$conf_from_proc" ] && [ -f "$conf_from_proc/elasticsearch.yml" ]; then
                            es_conf="$conf_from_proc/elasticsearch.yml"
                        fi
                    fi
                    if [ -z "$es_conf" ]; then
                        for f in /etc/elasticsearch/elasticsearch.yml /usr/local/etc/elasticsearch/elasticsearch.yml /usr/share/elasticsearch/config/elasticsearch.yml; do
                            if [ -f "$f" ]; then
                                es_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$es_conf" ]; then
                        printf '%s' "$es_conf"
                    fi
                    """,
                ),
                _path(
                    "ES_URL",
                    "Elasticsearch API URL",
                    "http://localhost:9200",
                    "원격 REST API 점검에 사용할 URL입니다. 탐지되지 않으면 수동 입력합니다.",
                    """
                    if command -v curl >/dev/null 2>&1; then
                        if curl -s -m 3 http://localhost:9200 >/dev/null 2>&1; then
                            printf '%s' http://localhost:9200
                        fi
                    fi
                    """,
                ),
            ),
        ),
        "MongoDB": ApplicationDefinition(
            key="MongoDB",
            display_name="MongoDB",
            script_path=SCRIPTS_DIR / "mongodb_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("mongosh", "command -v mongosh >/dev/null 2>&1 || command -v mongo >/dev/null 2>&1"),
                _rule("mongod", "ps -ef | grep -E '[m]ongod' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "MONGOD_CONF",
                    "MongoDB 설정 파일",
                    "/etc/mongod.conf",
                    "mongod --config 경로입니다.",
                    """
                    mongo_conf=''
                    mongod_proc=$(ps -ef 2>/dev/null | grep '[m]ongod' | grep -v mongos | head -1)
                    if [ -n "$mongod_proc" ]; then
                        conf_from_proc=$(printf '%s' "$mongod_proc" | sed -n 's/.*--config[= ]\\([^ ]*\\).*/\\1/p')
                        if [ -n "$conf_from_proc" ] && [ -f "$conf_from_proc" ]; then
                            mongo_conf="$conf_from_proc"
                        fi
                    fi
                    if [ -z "$mongo_conf" ]; then
                        for f in /etc/mongod.conf /etc/mongodb.conf /usr/local/etc/mongod.conf; do
                            if [ -f "$f" ]; then
                                mongo_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$mongo_conf" ]; then
                        printf '%s' "$mongo_conf"
                    fi
                    """,
                ),
            ),
        ),
        "Apache": ApplicationDefinition(
            key="Apache",
            display_name="Apache",
            script_path=SCRIPTS_DIR / "apache_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule(
                    "apache2",
                    "command -v apache2 >/dev/null 2>&1 || command -v httpd >/dev/null 2>&1 || command -v apachectl >/dev/null 2>&1 || test -x /usr/local/apache2/bin/httpd || test -x /usr/sbin/httpd || test -x /usr/sbin/apache2",
                ),
                _rule("apache-proc", "ps -ef | grep -E '[a]pache2|[h]ttpd' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "APACHE_CONF",
                    "Apache 설정 파일",
                    "/etc/apache2/apache2.conf",
                    "httpd.conf 또는 apache2.conf 전체 경로입니다.",
                    """
                    apache_conf=''
                    apache_bin=$(command -v httpd 2>/dev/null)
                    [ -z "$apache_bin" ] && apache_bin=$(command -v apache2 2>/dev/null)
                    [ -z "$apache_bin" ] && apache_bin=$(command -v apachectl 2>/dev/null)
                    if [ -n "$apache_bin" ]; then
                        server_root=$("$apache_bin" -V 2>/dev/null | sed -n 's/.*HTTPD_ROOT="\\(.*\\)"/\\1/p')
                        server_config=$("$apache_bin" -V 2>/dev/null | sed -n 's/.*SERVER_CONFIG_FILE="\\(.*\\)"/\\1/p')
                        if [ -n "$server_root" ] && [ -n "$server_config" ]; then
                            if printf '%s' "$server_config" | grep -q '^/'; then
                                apache_conf="$server_config"
                            else
                                apache_conf="$server_root/$server_config"
                            fi
                        fi
                    fi
                    if [ -z "$apache_conf" ]; then
                        for f in /etc/httpd/conf/httpd.conf /etc/apache2/apache2.conf /usr/local/apache2/conf/httpd.conf; do
                            if [ -f "$f" ]; then
                                apache_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$apache_conf" ]; then
                        printf '%s' "$apache_conf"
                    fi
                    """,
                ),
            ),
        ),
        "Nginx": ApplicationDefinition(
            key="Nginx",
            display_name="Nginx",
            script_path=SCRIPTS_DIR / "nginx_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("nginx", "command -v nginx >/dev/null 2>&1"),
                _rule("nginx-proc", "ps -ef | grep -E '[n]ginx' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "NGINX_CONF",
                    "Nginx 설정 파일",
                    "/etc/nginx/nginx.conf",
                    "nginx.conf 전체 경로입니다.",
                    """
                    nginx_conf=''
                    nginx_bin=$(command -v nginx 2>/dev/null)
                    if [ -n "$nginx_bin" ]; then
                        nginx_test=$("$nginx_bin" -t 2>&1)
                        conf_from_test=$(printf '%s' "$nginx_test" | sed -n 's/.*configuration file \\(.*\\) test.*/\\1/p')
                        if [ -n "$conf_from_test" ] && [ -f "$conf_from_test" ]; then
                            nginx_conf="$conf_from_test"
                        fi
                    fi
                    if [ -z "$nginx_conf" ]; then
                        for f in /etc/nginx/nginx.conf /usr/local/nginx/conf/nginx.conf /usr/local/etc/nginx/nginx.conf; do
                            if [ -f "$f" ]; then
                                nginx_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$nginx_conf" ]; then
                        printf '%s' "$nginx_conf"
                    fi
                    """,
                ),
            ),
        ),
        "Tomcat": ApplicationDefinition(
            key="Tomcat",
            display_name="Tomcat",
            script_path=SCRIPTS_DIR / "tomcat_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("catalina", "ps -ef | grep -E '[c]atalina|[t]omcat' >/dev/null 2>&1"),
                _rule("tomcat-dir", "ls /opt/tomcat /usr/share/tomcat >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "CATALINA_HOME",
                    "Tomcat 홈 디렉터리",
                    "/usr/local/tomcat",
                    "conf/server.xml 이 있는 Tomcat 홈입니다.",
                    """
                    if [ -n "$CATALINA_HOME" ] && [ -d "$CATALINA_HOME" ]; then
                        printf '%s' "$CATALINA_HOME"
                    else
                        tomcat_proc=$(ps -ef 2>/dev/null | grep -E '[c]atalina|[t]omcat' | head -1)
                        if [ -n "$tomcat_proc" ]; then
                            home_from_proc=$(printf '%s' "$tomcat_proc" | sed -n 's/.*-Dcatalina\\.home=\\([^ ]*\\).*/\\1/p')
                            if [ -n "$home_from_proc" ] && [ -d "$home_from_proc" ]; then
                                printf '%s' "$home_from_proc"
                            fi
                        fi
                        if [ -z "$home_from_proc" ]; then
                            for d in /usr/share/tomcat* /opt/tomcat* /var/lib/tomcat* /usr/local/tomcat*; do
                                if [ -d "$d" ] && [ -f "$d/conf/server.xml" ]; then
                                    printf '%s' "$d"
                                    break
                                fi
                            done
                        fi
                    fi
                    """,
                ),
            ),
        ),
        "Docker": ApplicationDefinition(
            key="Docker",
            display_name="Docker",
            script_path=SCRIPTS_DIR / "docker_cce_check.sh",
            shell="sh",
            shell_family="posix",
            probe_rules=(
                _rule("docker", "command -v docker >/dev/null 2>&1"),
                _rule("dockerd", "ps -ef | grep -E '[d]ockerd' >/dev/null 2>&1"),
                _rule("docker-sock", "test -S /var/run/docker.sock"),
            ),
            path_settings=(
                _path(
                    "DOCKER_CONF",
                    "Docker daemon 설정 파일",
                    "/etc/docker/daemon.json",
                    "dockerd --config-file 또는 daemon.json 경로입니다.",
                    """
                    docker_conf=''
                    dockerd_proc=$(ps -ef 2>/dev/null | grep '[d]ockerd' | head -1)
                    if [ -n "$dockerd_proc" ]; then
                        conf_from_proc=$(printf '%s' "$dockerd_proc" | sed -n 's/.*--config-file[= ]\\([^ ]*\\).*/\\1/p')
                        if [ -n "$conf_from_proc" ] && [ -f "$conf_from_proc" ]; then
                            docker_conf="$conf_from_proc"
                        fi
                    fi
                    if [ -z "$docker_conf" ]; then
                        for f in /etc/docker/daemon.json "$HOME/.docker/daemon.json"; do
                            if [ -f "$f" ]; then
                                docker_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$docker_conf" ]; then
                        printf '%s' "$docker_conf"
                    fi
                    """,
                ),
            ),
        ),
        "K8s(Master)": ApplicationDefinition(
            key="K8s(Master)",
            display_name="Kubernetes Master",
            script_path=SCRIPTS_DIR / "k8s_master_cce_check.sh",
            shell="sh",
            shell_family="posix",
            probe_rules=(
                _rule("kubectl", "command -v kubectl >/dev/null 2>&1"),
                _rule("kube-apiserver", "ps -ef | grep -E '[k]ube-apiserver' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "K8S_MANIFEST_DIR",
                    "Kubernetes 매니페스트 디렉터리",
                    "/etc/kubernetes/manifests",
                    "정적 Pod 매니페스트 디렉터리입니다.",
                    """
                    apiserver_proc=$(ps -ef 2>/dev/null | grep '[k]ube-apiserver' | head -1)
                    manifest_dir=$(printf '%s' "$apiserver_proc" | sed -n 's/.*--pod-manifest-path[= ]\\([^ ]*\\).*/\\1/p')
                    if [ -n "$manifest_dir" ] && [ -d "$manifest_dir" ]; then
                        printf '%s' "$manifest_dir"
                    else
                        for d in /etc/kubernetes/manifests /etc/kubernetes; do
                            if [ -d "$d" ]; then
                                printf '%s' "$d"
                                break
                            fi
                        done
                    fi
                    """,
                ),
            ),
        ),
        "K8s(Worker)": ApplicationDefinition(
            key="K8s(Worker)",
            display_name="Kubernetes Worker",
            script_path=SCRIPTS_DIR / "k8s_worker_cce_check.sh",
            shell="sh",
            shell_family="posix",
            probe_rules=(
                _rule("kubelet", "ps -ef | grep -E '[k]ubelet' >/dev/null 2>&1"),
                _rule("kubeconfig", "test -f /etc/kubernetes/kubelet.conf"),
                _rule("k3s-agent", "ps -ef | grep -E '[k]3s agent' >/dev/null 2>&1"),
                _rule("k3s-agent-dir", "test -d /var/lib/rancher/k3s/agent"),
            ),
            path_settings=(
                _path(
                    "KUBELET_CONF",
                    "kubelet 설정 파일",
                    "/var/lib/kubelet/config.yaml",
                    "kubelet --config 경로입니다.",
                    """
                    kubelet_conf=''
                    kubelet_proc=$(ps -ef 2>/dev/null | grep '[k]ubelet' | head -1)
                    if [ -n "$kubelet_proc" ]; then
                        conf_from_proc=$(printf '%s' "$kubelet_proc" | sed -n 's/.*--config[= ]\\([^ ]*\\).*/\\1/p')
                        if [ -n "$conf_from_proc" ] && [ -f "$conf_from_proc" ]; then
                            kubelet_conf="$conf_from_proc"
                        fi
                    fi
                    if [ -z "$kubelet_conf" ]; then
                        for f in /var/lib/kubelet/config.yaml /etc/kubernetes/kubelet.conf /var/lib/rancher/k3s/agent/kubelet.kubeconfig; do
                            if [ -f "$f" ]; then
                                kubelet_conf="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$kubelet_conf" ]; then
                        printf '%s' "$kubelet_conf"
                    fi
                    """,
                ),
            ),
        ),
        "PHP": ApplicationDefinition(
            key="PHP",
            display_name="PHP",
            script_path=SCRIPTS_DIR / "php_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("php", "command -v php >/dev/null 2>&1"),
                _rule("php-fpm", "ps -ef | grep -E '[p]hp-fpm' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "PHP_INI",
                    "PHP ini 파일",
                    "/etc/php/8.3/cli/php.ini",
                    "php --ini 결과의 Loaded Configuration File 경로입니다.",
                    """
                    php_ini=''
                    php_bin=$(command -v php 2>/dev/null)
                    if [ -n "$php_bin" ]; then
                        php_ini=$("$php_bin" --ini 2>/dev/null | sed -n 's/.*Loaded Configuration File:[[:space:]]*\\(.*\\)/\\1/p')
                        if [ "$php_ini" = "(none)" ]; then
                            php_ini=''
                        fi
                    fi
                    if [ -z "$php_ini" ]; then
                        for f in /etc/php/*/cli/php.ini /etc/php/*/fpm/php.ini /etc/php.ini /usr/local/etc/php/php.ini; do
                            if [ -f "$f" ]; then
                                php_ini="$f"
                                break
                            fi
                        done
                    fi
                    if [ -n "$php_ini" ]; then
                        printf '%s' "$php_ini"
                    fi
                    """,
                ),
            ),
        ),
        "NodeJS": ApplicationDefinition(
            key="NodeJS",
            display_name="Node.js",
            script_path=SCRIPTS_DIR / "nodejs_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("node", "command -v node >/dev/null 2>&1"),
                _rule("node-proc", "ps -ef | grep -E '[n]ode' >/dev/null 2>&1"),
            ),
        ),
        "Hadoop": ApplicationDefinition(
            key="Hadoop",
            display_name="Hadoop",
            script_path=SCRIPTS_DIR / "hadoop_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("hdfs", "command -v hdfs >/dev/null 2>&1"),
                _rule("namenode", "ps -ef | grep -E '[N]ameNode|[D]ataNode|[R]esourceManager|[N]odeManager' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "HADOOP_HOME",
                    "Hadoop 홈 디렉터리",
                    "/opt/hadoop",
                    "bin/hadoop 가 위치한 홈 디렉터리입니다.",
                    """
                    if [ -n "$HADOOP_HOME" ] && [ -d "$HADOOP_HOME" ]; then
                        printf '%s' "$HADOOP_HOME"
                    else
                        hadoop_bin=$(command -v hadoop 2>/dev/null)
                        if [ -n "$hadoop_bin" ]; then
                            printf '%s' "$(cd "$(dirname "$hadoop_bin")/.." && pwd)"
                        fi
                    fi
                    """,
                ),
                _path(
                    "HADOOP_CONF_DIR",
                    "Hadoop 설정 디렉터리",
                    "/etc/hadoop/conf",
                    "core-site.xml 등이 있는 설정 디렉터리입니다.",
                    """
                    if [ -n "$HADOOP_CONF_DIR" ] && [ -d "$HADOOP_CONF_DIR" ]; then
                        printf '%s' "$HADOOP_CONF_DIR"
                    else
                        for d in /etc/hadoop/conf /opt/hadoop*/etc/hadoop /usr/lib/hadoop/etc/hadoop /usr/local/hadoop/etc/hadoop; do
                            if [ -d "$d" ]; then
                                printf '%s' "$d"
                                break
                            fi
                        done
                    fi
                    """,
                ),
            ),
        ),
        "Ceph": ApplicationDefinition(
            key="Ceph",
            display_name="Ceph",
            script_path=SCRIPTS_DIR / "ceph_cce_check.sh",
            shell="bash",
            shell_family="posix",
            probe_rules=(
                _rule("ceph", "command -v ceph >/dev/null 2>&1"),
                _rule("ceph-proc", "ps -ef | grep -E '[c]eph-mon|[c]eph-osd|[c]eph-mgr' >/dev/null 2>&1"),
            ),
            path_settings=(
                _path(
                    "CEPH_CONF",
                    "Ceph 설정 파일",
                    "/etc/ceph/ceph.conf",
                    "ceph.conf 전체 경로입니다.",
                    """
                    if [ -n "$CEPH_CONF" ] && [ -f "$CEPH_CONF" ]; then
                        printf '%s' "$CEPH_CONF"
                    else
                        for f in /etc/ceph/ceph.conf /usr/local/etc/ceph/ceph.conf; do
                            if [ -f "$f" ]; then
                                printf '%s' "$f"
                                break
                            fi
                        done
                    fi
                    """,
                ),
            ),
        ),
    }
)

ACTIVE_APP_KEYS = tuple(
    app_key for app_key, app in APP_DEFINITIONS.items() if app.shell_family == "posix"
)


FRAMEWORK_HEURISTICS = {
    "클라우드": ["공공CSAP(초기매핑)"],
    "기반시설": ["ISMS-P(초기매핑)"],
    "통합": ["공공CSAP(초기매핑)", "ISMS-P(초기매핑)"],
    "공통": ["공공CSAP(초기매핑)", "ISMS-P(초기매핑)"],
}


def derive_framework_tags(source: str) -> list[str]:
    return FRAMEWORK_HEURISTICS.get(source, [])


@lru_cache(maxsize=1)
def load_control_catalog() -> dict:
    workbook = load_workbook(MERGED_ITEMS_XLSX, read_only=True, data_only=True)
    sheet = workbook["CCE 항목 통합(중복제거)"]
    by_app: dict[str, list[dict]] = defaultdict(list)
    by_code: dict[str, dict] = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        target, source, importance, code, title, category, diagnosis, remediation = row[:8]
        if not target or not code:
            continue
        if str(target).startswith("  "):
            continue

        raw_code = str(code).strip()
        control = {
            "target": str(target).strip(),
            "source": (source or "").strip(),
            "importance": (importance or "").strip(),
            "code": to_preferred_code_text(raw_code),
            "title": (title or "").strip(),
            "category": (category or "").strip(),
            "diagnosis": (diagnosis or "").strip(),
            "remediation": (remediation or "").strip(),
            "frameworks": derive_framework_tags((source or "").strip()),
        }
        by_app[control["target"]].append(control)
        for alias in control_code_aliases(raw_code) | control_code_aliases(control["code"]):
            by_code[f'{control["target"]}:{alias}'] = control

    workbook.close()
    return {
        "by_app": {key: list(value) for key, value in by_app.items()},
        "by_code": by_code,
    }


def get_control(app_key: str, code: str) -> dict | None:
    return load_control_catalog()["by_code"].get(f"{app_key}:{code}")


def get_controls_for_app(app_key: str) -> list[dict]:
    return load_control_catalog()["by_app"].get(app_key, [])


def list_controls(app_key: str | None = None) -> list[dict]:
    catalog = load_control_catalog()["by_app"]
    if app_key:
        return list(catalog.get(app_key, []))
    controls: list[dict] = []
    for value in catalog.values():
        controls.extend(value)
    return controls


def get_app_definition(app_key: str) -> ApplicationDefinition | None:
    return APP_DEFINITIONS.get(app_key)


def list_app_definitions() -> list[ApplicationDefinition]:
    return [APP_DEFINITIONS[key] for key in ACTIVE_APP_KEYS]


def list_all_app_definitions() -> list[ApplicationDefinition]:
    return list(APP_DEFINITIONS.values())


def host_support_for_app(shell_type: str, transport: str, app_key: str) -> dict:
    app = get_app_definition(app_key)
    if app is None:
        return {"supported": False, "reason": "정의되지 않은 애플리케이션"}
    if app.shell_family != "posix":
        return {"supported": False, "reason": "Windows/PowerShell 경로는 후속 단계에서 지원 예정"}
    if shell_type != app.shell_family:
        return {
            "supported": False,
            "reason": f"호스트 셸({shell_type})과 스크립트 셸({app.shell_family})이 일치하지 않음",
        }
    if transport not in {"local", "ssh"}:
        if transport == "compose":
            return {"supported": True, "reason": "Docker Compose 테스트 랩 컨테이너에서 실행 가능"}
        return {"supported": False, "reason": f"지원하지 않는 전송방식: {transport}"}
    if shell_type == "powershell":
        return {"supported": False, "reason": "PowerShell 원격 실행 어댑터는 다음 단계 구현 대상"}
    return {"supported": True, "reason": "현재 하네스에서 실행 가능"}


def framework_options() -> list[str]:
    values = OrderedDict()
    for control in list_controls():
        for framework in control["frameworks"]:
            values[framework] = None
    return list(values.keys())


def serialize_frameworks(frameworks: list[str]) -> str:
    return json.dumps(frameworks, ensure_ascii=False)


def deserialize_frameworks(raw_value: str | None) -> list[str]:
    if not raw_value:
        return []
    return json.loads(raw_value)
